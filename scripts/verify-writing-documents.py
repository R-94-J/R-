"""校验写作规范原件哈希及开题报告PDF。需安装python-docx、pymupdf。"""
import hashlib
import json
import re
from pathlib import Path

import pymupdf
from docx import Document
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
STANDARDS = ROOT / 'docs/standards'
PROPOSAL = ROOT / 'docs/proposal'


def main():
    manifest = json.loads((STANDARDS / 'download-manifest.json').read_text(encoding='utf-8'))
    attachments = []
    for item in manifest:
        path = STANDARDS / item['file']
        data = path.read_bytes()
        assert len(data) == item['bytes'], path
        assert hashlib.sha256(data).hexdigest() == item['sha256'], path
        record = {'file': item['file'], 'sha256Valid': True, 'bytes': len(data)}
        if path.suffix == '.pdf':
            assert data.startswith(b'%PDF'), path
            with pymupdf.open(path) as pdf:
                record['pages'] = len(pdf)
                assert len(pdf) >= 5, path
        attachments.append(record)
    assert sum(item['file'].endswith('.pdf') for item in attachments) == 5
    document = Document(PROPOSAL / '开题报告.docx')
    table = document.tables[0]
    assert len(table.rows) == 13
    assert document.sections[0].footer._element.xpath('.//w:fldSimple[@w:instr="PAGE"]')
    for row in range(5, 10):
        assert table.cell(row, 0).text.startswith(f'{row - 4}　')
    refs = table.cell(10, 1).paragraphs
    assert [int(re.match(r'\[(\d+)\]', p.text).group(1)) for p in refs] == list(range(1, 12))
    for p in refs:
        assert not any(r.font.superscript for r in p.runs)
    for row in range(5, 10):
        for p in table.cell(row, 1).paragraphs:
            for run in p.runs:
                if re.fullmatch(r'\[\d+(?:[-,]\d+)*\]', run.text):
                    assert run.font.superscript
    with pymupdf.open(PROPOSAL / '开题报告.pdf') as pdf:
        text = '\n'.join(page.get_text(sort=True) for page in pdf)
        pages = []
        fonts = set()
        for i, page in enumerate(pdf):
            footer = page.get_text(clip=pymupdf.Rect(0, page.rect.height - 80,
                                                    page.rect.width, page.rect.height)).strip()
            assert footer == str(i + 1), (i + 1, footer)
            assert abs(page.rect.width - 595.3) < 1 and abs(page.rect.height - 841.9) < 1
            pages.append({'page': i + 1, 'sizePt': [page.rect.width, page.rect.height], 'footer': footer})
            fonts.update(font[3] for font in page.get_fonts())
        assert all(word not in text for word in ['西班牙', '迟苏丹', '訾月琳', '通过审核'])
        report = {'pages': len(pdf), 'pageDetails': pages, 'fonts': sorted(fonts),
                  'pageNumbersContinuous': True, 'noOldExample': True,
                  'allElevenReferencesPresent': all(f'[{i}]' in text for i in range(1, 12)),
                  'schoolWritingNormFullTextVerified': False}
        assert report['allElevenReferencesPresent']
        (PROPOSAL / 'pdf-verification.json').write_text(
            json.dumps(report, ensure_ascii=False, indent=2) + '\n', encoding='utf-8')
    result = {'verifiedAt': '2026-10-16', 'requestedDocuments': 6, 'standardPdfsObtained': 5,
              'schoolWritingNormFullTextObtained': False, 'attachments': attachments,
              'reportPdfPages': report['pages'], 'scope': '文件完整性及确定性格式检查，不是查重或AI检测'}
    (STANDARDS / 'verification.json').write_text(
        json.dumps(result, ensure_ascii=False, indent=2) + '\n', encoding='utf-8')
    print(json.dumps(result, ensure_ascii=False, indent=2))


if __name__ == '__main__':
    main()
