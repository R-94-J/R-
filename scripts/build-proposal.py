"""把开题报告正文填入教务处原表，保留表格结构和页面设置。依赖 python-docx。"""
import copy
import hashlib
import json
import re
from pathlib import Path
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.shared import Pt

ROOT = Path(__file__).resolve().parents[1]
FOLDER = ROOT / 'docs/proposal'
TEMPLATE = FOLDER / 'sources/开题报告空白模板.docx'
TITLE = '基于检索增强生成的校园知识库问答系统设计与实现'
NS = {'w': 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}


def clear_cell(cell):
    for child in list(cell._tc):
        if child.tag != qn('w:tcPr'):
            cell._tc.remove(child)


def paragraph(container, text, *, bold=False, center=False, indent=False, size=12, citations=False):
    p = container.add_paragraph()
    p.style = 'Normal'
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER if center else WD_ALIGN_PARAGRAPH.LEFT if bold else WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.first_line_indent = Pt(24 if indent else 0)
    p.paragraph_format.keep_together = False
    p.paragraph_format.keep_with_next = bold
    p.paragraph_format.widow_control = True
    parts = re.split(r'(\[\d+(?:[-,]\d+)*\])', text) if citations else [text]
    for part in parts:
        if not part:
            continue
        run = p.add_run(part)
        run.font.size = Pt(size)
        run.bold = bold
        if citations and re.fullmatch(r'\[\d+(?:[-,]\d+)*\]', part):
            run.font.superscript = True
        fonts = run._r.get_or_add_rPr().get_or_add_rFonts()
        for key,value in [('eastAsia','宋体'),('ascii','Times New Roman'),('hAnsi','Times New Roman')]:
            fonts.set(qn('w:' + key),value)
    return p


def set_cell(cell, text, center=True):
    clear_cell(cell)
    paragraph(cell,text,center=center)


def prepare_template():
    if TEMPLATE.exists():
        return
    extracted = FOLDER / 'sources/毕业论文（设计）表格模板'
    original = next(extracted.glob('04*.docx'))
    doc = Document(original)
    table = doc.tables[0]
    # 官网压缩包中的原表示例已填写个人信息和审核意见，全部清空。
    for row,col in [(0,1),(0,5),(1,1),(1,5),(2,1),(2,3),(2,6),(3,1)]:
        set_cell(table.cell(row,col),'')
    for row in range(5,13):
        set_cell(table.cell(row,1),'',center=False)
    p = doc.paragraphs[2]
    p.clear()
    r=p.add_run('学院：________________              填表日期：____年__月__日')
    r.font.size=Pt(12)
    # 移除示例签名图片关系，保证交付文件没有遗留的签名媒体。
    for key,rel in list(doc.part.rels.items()):
        if rel.reltype.endswith('/image'):
            doc.part.drop_rel(key)
    doc.core_properties.author = ''
    doc.core_properties.last_modified_by = ''
    doc.core_properties.title = '山东青年政治学院毕业论文（设计）开题报告空白表'
    doc.save(TEMPLATE)
    record={'downloadedTemplate':original.name,'sha256':hashlib.sha256(original.read_bytes()).hexdigest(),
            'sourceArchive':'official-forms.rar','cleanedTemplate':TEMPLATE.name,
            'changes':['清空示例个人信息、正文、审核意见、日期及签名图片','保留原表标题、标签、合并单元格和节设置']}
    (FOLDER/'sources/template-record.json').write_text(json.dumps(record,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')


def main():
    prepare_template()
    doc = Document(TEMPLATE)
    original_section = copy.deepcopy(doc.sections[0]._sectPr)
    original_grid = copy.deepcopy(doc.tables[0]._tbl.tblGrid)
    table=doc.tables[0]
    doc.paragraphs[2].clear().add_run('学院：职业教育学院              填表日期：____年__月__日').font.size=Pt(12)
    for row,col,text in [(0,1,'待填写'),(0,5,'待填写'),(1,1,'职业教育学院'),(1,5,'大数据工程技术'),
                         (2,1,''),(2,3,''),(2,6,''),(3,1,TITLE)]:
        set_cell(table.cell(row,col),text)
    content=(FOLDER/'开题报告.md').read_text(encoding='utf-8')
    sections={}
    current=None
    for line in content.splitlines():
        if line.startswith('## '):
            current=re.sub(r'^\d+[\s　]+', '', line[3:]); sections[current]=[]
        elif current and line.strip():
            sections[current].append(line.strip())
    mapping={5:'选题依据及研究内容',6:'理论和实践意义',7:'论文撰写过程中拟采取的方法和手段',
             8:'论文撰写提纲',9:'计划进度及其内容',10:'参考文献'}
    # 原表栏目顺序、合并关系和边框不动，正文栏目统一采用阿拉伯数字章序。
    for number, row in enumerate(range(5, 10), start=1):
        label = table.cell(row, 0)
        original_note = label.paragraphs[1].text if row == 5 and len(label.paragraphs) > 1 else ''
        clear_cell(label)
        label.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
        p = paragraph(label, f'{number}　{mapping[row]}', bold=True)
        p.paragraph_format.keep_with_next = False
        p.paragraph_format.widow_control = False
        if original_note:
            p = paragraph(label, original_note)
            p.paragraph_format.widow_control = False
    # 正文引用必须对应文末条目，序号连续；避免修改稿漏删引用或文献。
    references = [int(re.match(r'\[(\d+)\]', line).group(1)) for line in sections['参考文献']]
    assert references == list(range(1, len(references) + 1)), '参考文献编号不连续'
    cited = set()
    first_citations = []
    body = content.split('## 参考文献')[0]
    for start, end in re.findall(r'\[(\d+)(?:-(\d+))?\]', body):
        for number in range(int(start), int(end or start) + 1):
            if number not in cited:
                first_citations.append(number)
            cited.add(number)
    assert cited == set(references), '正文引用和参考文献不对应'
    assert first_citations == references, '参考文献序号与正文首次引用顺序不一致'
    for row,heading in mapping.items():
        cell=table.cell(row,1); clear_cell(cell)
        for line in sections[heading]:
            heading=line.startswith('### ')
            outline_title = row == 8 and re.match(r'^\d+　', line)
            bold = heading or bool(outline_title)
            p = paragraph(cell,line[4:] if heading else line,bold=bold,
                          indent=not bold and row in [5,6,7],citations=row!=10)
            if row == 10:
                # 文献采用悬挂缩进；列表序号保持基线，正文引文为上标。
                p.paragraph_format.left_indent = Pt(24)
                p.paragraph_format.first_line_indent = Pt(-24)
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                p.paragraph_format.keep_together = True
    set_cell(table.cell(11,1),'\n\n\n是否同意开题：________（是、否）\n\n指导教师（签字）：________  ____年__月__日',center=False)
    set_cell(table.cell(12,1),'\n\n\n\n分管副院长（签字）：________  ____年__月__日',center=False)
    # 两个审批区集中在同页，避免仅一个签字栏落在末页。
    for row_index in [11, 12]:
        properties = table.rows[row_index]._tr.get_or_add_trPr()
        if properties.find(qn('w:cantSplit')) is None:
            properties.append(OxmlElement('w:cantSplit'))
        for cell in table.rows[row_index].cells:
            for p in cell.paragraphs:
                p.paragraph_format.keep_with_next = row_index == 11
    doc.core_properties.author=''
    doc.core_properties.last_modified_by=''
    doc.core_properties.title=TITLE+'（开题报告）'
    # 开题表无独立前置页，整份报告使用连续阿拉伯页码。
    footer = doc.sections[0].footer
    p = footer.paragraphs[0]
    p.clear()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    field = OxmlElement('w:fldSimple')
    field.set(qn('w:instr'), 'PAGE')
    number_run = OxmlElement('w:r')
    number_props = OxmlElement('w:rPr')
    number_size = OxmlElement('w:sz')
    number_size.set(qn('w:val'), '21')
    number_props.append(number_size)
    number_run.append(number_props)
    number_text = OxmlElement('w:t')
    number_text.text = '1'
    number_run.append(number_text)
    field.append(number_run)
    p._p.append(field)
    output=FOLDER/'开题报告.docx'
    doc.save(output)
    # 结构、字号和残留示例的自动核对。
    verified=Document(output)
    def structure(element):
        return (element.tag,dict(element.attrib),[structure(child) for child in element if child.tag != qn('w:footerReference')])
    assert structure(verified.sections[0]._sectPr) == structure(original_section)
    assert structure(verified.tables[0]._tbl.tblGrid) == structure(original_grid)
    assert len(verified.tables[0].rows)==13
    text='\n'.join(p.text for p in verified.paragraphs)+'\n'+'\n'.join(
        c.text for row in verified.tables[0].rows for c in row.cells)
    for unwanted in ['西班牙','迟苏丹','訾月琳','通过审核','2017年']:
        assert unwanted not in text,unwanted
    assert not verified.part.element.xpath('//w:drawing')
    assert verified.sections[0].footer._element.xpath('.//w:fldSimple[@w:instr="PAGE"]')
    citation_runs = 0
    for row in mapping:
        for p in verified.tables[0].cell(row,1).paragraphs:
            for run in p.runs:
                assert run.font.size == Pt(12)
                assert run._r.rPr.rFonts.get(qn('w:eastAsia'))=='宋体'
                if row != 10 and re.fullmatch(r'\[\d+(?:[-,]\d+)*\]', run.text):
                    assert run.font.superscript is True
                    citation_runs += 1
                if row == 10:
                    assert run.font.superscript is not True
    assert citation_runs > 0
    headings = [line[4:] for lines in sections.values() for line in lines if line.startswith('### ')]
    assert [heading.split('　')[0] for heading in headings] == [
        '1.1','1.2','1.3','1.4','2.1','2.2','3.1','3.2','3.3','3.4']
    report={'file':output.name,'template':'sources/'+TEMPLATE.name,'rows':13,
            'paper':'A4','marginsCm':{'top':2.54,'bottom':2.54,'left':3.175,'right':3.175},
            'titlePt':16,'bodyPt':12,'chineseFont':'宋体','latinFont':'Times New Roman',
            'contentChineseCharacters':len(re.findall('[\u4e00-\u9fff]',content)),
            'personalFields':'待填写','reviewComments':'留空','signatureImages':0,
            'referenceCount':len(references),'allReferencesCited':True,
            'pageLayoutAndTableGridMatchTemplate':True,'bodyNumbering':'1, 1.1, 1.2…',
            'citationStyle':'顺序编码制，方括号上标','citationRuns':citation_runs,'pageNumberField':'PAGE',
            'schoolWritingNormFullTextVerified':False}
    (FOLDER/'format-verification.json').write_text(json.dumps(report,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
    print(json.dumps(report,ensure_ascii=False,indent=2))


if __name__=='__main__':
    main()
